import concurrent.futures
import random
import string
import pandas as pd
import os
import time
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
import threading
from pyfiglet import Figlet
from tkinter import Tk, filedialog
import openpyxl

# Function to get the date from a file
def get_date_from_file(filename='date.txt'):
    try:
        with open(filename, 'r') as file:
            date = file.read().strip()
        return date
    except Exception as e:
        raise EmailSendingError(f"Error reading date from file: {str(e)}")

# Define a custom exception class for email sending errors
class EmailSendingError(Exception):
    pass

# Function to generate a random alphanumeric string
def generate_string(length=7):
    alphanumeric = ''.join(random.choice(string.ascii_uppercase + string.digits) for _ in range(length))
    return alphanumeric

# Function to generate a random 8-digit numeric string
def generate_random_numeric():
    random_number = ''.join([str(random.randint(0, 9)) for _ in range(8)])
    return random_number

# Function to read recipients from a file
def read_recipients(filename='recipients.xlsx'):
    recipients_df = pd.read_excel(filename)
    return recipients_df[['Customer Name', 'Email']].values.tolist()

# Function to read sender details from a file
def read_senders(filename='senders.xlsx'):
    senders_df = pd.read_excel(filename)
    required_columns = ['Email', 'Password', 'SenderName', 'SMTP', 'Port']
    if not all(column in senders_df.columns for column in required_columns):
        print("Error: The 'senders.xlsx' file is missing required columns.")
        exit()
    return senders_df.to_dict(orient='records')

# Function to read subject lines from a file and shuffle them
def read_subject_lines(filename='subjects.xlsx'):
    subject_lines_df = pd.read_excel(filename)
    subject_lines = subject_lines_df['Subject'].tolist()
    random.shuffle(subject_lines)
    return subject_lines

# Function to get the HTML file using a file dialog
def get_html_file():
    root = Tk()
    root.withdraw()  # Hide the main window
    file_path = filedialog.askopenfilename(filetypes=[("HTML Files", "*.html")])
    return file_path

# Function to load HTML content from a file
def load_html_content(html_file):
    with open(html_file, 'r', encoding='utf-8') as file:
        return file.read()

# Replace placeholders in the HTML content
def replace_placeholders(html_content, amount, phone, random_numeric, name, date):
    html_content = html_content.replace('{{amount}}', str(amount))
    html_content = html_content.replace('{{phone}}', str(phone))
    html_content = html_content.replace('{{name}}', str(name))
    html_content = html_content.replace('{{date}}', date)  # Add date replacement
    html_content = html_content.replace('{{RAN}}', str(random_numeric))
    return html_content

# Function to send an email
def send_email(sender, sender_password, sender_name, recipient, subject, message):
    try:
        with smtplib.SMTP(sender['SMTP'], sender['Port']) as server:
            server.starttls()
            server.login(sender['Email'], sender_password)

            msg = MIMEMultipart()
            msg['From'] = f"{sender_name} <{sender['Email']}>"
            msg['To'] = recipient[1]
            msg['Subject'] = subject

            msg.attach(MIMEText(message, "html"))

            server.sendmail(sender['Email'], recipient[1], msg.as_string())

            print("\033[92m" + f'Successfully sent email from {sender["Email"]} to {recipient[1]}' + "\033[0m")

            return sender['Email']

    except Exception as e:
        error_message = str(e)
        if "'float' object has no attribute 'encode'" in error_message:
            raise EmailSendingError("Error sending email: 'float' object has no attribute 'encode'")
        else:
            print("\033[91m" + f'Error sending email from {sender["Email"]} to {recipient[1]}: {error_message}' + "\033[0m")
            return None

# Function to select a random 'amount' from the 'data.xlsx' file
def select_random_amount(filename='data.xlsx'):
    data_df = pd.read_excel(filename)
    random_amount = data_df.sample(n=1)['amount'].values[0]
    return random_amount

# Function to select a random name from the 'data.xlsx' file
def select_random_name(filename='data.xlsx'):
    data_df = pd.read_excel(filename)
    random_name = data_df.sample(n=1)['name'].values[0]
    return random_name

# Function to always select 'phone' from cell B2 in the 'data.xlsx' file
def select_fixed_phone(filename='data.xlsx'):
    try:
        workbook = openpyxl.load_workbook(filename)
        sheet = workbook.active
        fixed_phone = sheet['B2'].value
        return fixed_phone
    except Exception as e:
        print("Error reading phone from Excel:", str(e))
        return None

# Function to send emails concurrently
def send_emails_concurrently(senders, recipients, subject_lines, html_content, random_numeric, delay):
    sender_count = len(senders)

    if sender_count == 0:
        print("No sender accounts available. Exiting.")
        return

    recipient_count = len(recipients)

    if recipient_count == 0:
        print("No recipients available. Exiting.")
        return

    unsent_recipients = []
    sent_counts = []
    senders_lock = threading.Lock()

    def send_email_task(sender, sender_password, sender_name, recipient, subject, message):
        result = send_email(sender, sender_password, sender_name, recipient, subject, message)
        if result is not None:
            with senders_lock:
                sent_counts.append(result)

    with concurrent.futures.ThreadPoolExecutor(max_workers=sender_count) as executor:
        futures = []

        for i in range(recipient_count):
            recipient = recipients[i]
            recipient_name, recipient_email = recipient

            sender = random.choice(senders)
            sender_index = senders.index(sender)
            sender_email = sender['Email']
            sender_password = sender['Password']
            sender_name = sender['SenderName']

            subject = random.choice(subject_lines)
            subject += f"{alphanumeric_string}"

            # Get 'amount' and 'phone'
            amount = select_random_amount()
            phone = select_fixed_phone()
            name = select_random_name()

            # Get the date from date.txt
            date = get_date_from_file()

            # Replace placeholders in the HTML content
            modified_html_content = replace_placeholders(html_content, amount, phone, random_numeric, name, date)

            # Send the email (concurrently)
            future = executor.submit(send_email_task, sender, sender_password, sender_name, recipient, subject, modified_html_content)
            futures.append(future)

            if unsent_recipients:
                with open('unsent.txt', 'w') as file:
                    for recipient in unsent_recipients:
                        file.write(f"{recipient}\n")

            time.sleep(delay)

        concurrent.futures.wait(futures)

if __name__ == "__main__":
    try:
        sent_counts = []  # Initialize the sent_counts list
        alphanumeric_string = generate_string()
        random_numeric = generate_random_numeric()
        recipients = read_recipients('recipients.xlsx')

        if not recipients:
            print("No recipients available. Exiting.")
            exit()

        senders = read_senders('senders.xlsx')
        subject_lines = read_subject_lines('subjects.xlsx')
        html_file = get_html_file()
        html_content = load_html_content(html_file)
        delay = float(input("Enter the delay between each email (in seconds): "))

        send_emails_concurrently(senders, recipients, subject_lines, html_content, random_numeric, delay)

        for sender in sent_counts:
            print(f"Email sent using sender: {sender}")
    except EmailSendingError as e:
        print(f"Email sending process stopped due to an error: {e}")
    except Exception as e:
        print(f"An unexpected error occurred: {e}")
